import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins: 1 inch all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    styles = doc.styles

    # Base Normal Font
    style_normal = styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Segoe UI'
    font_normal.size = Pt(10.5)
    font_normal.color.rgb = RGBColor(51, 51, 51) # #333333

    # Primary Colors
    PRIMARY = RGBColor(37, 99, 235)      # Royal Blue (#2563EB)
    SECONDARY = RGBColor(30, 41, 59)    # Slate Dark (#1E293B)
    MUTED = RGBColor(100, 116, 139)     # Slate Muted (#64748B)

    # Title Page / Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    run_icon = p_title.add_run("🍽️ ")
    run_icon.font.size = Pt(28)
    run_title = p_title.add_run("CanteenQ")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(32)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Smart Campus Canteen Queue & Food Ordering System")
    run_sub.font.name = 'Segoe UI Semibold'
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = SECONDARY

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_after = Pt(36)
    run_tag = p_tag.add_run("Comprehensive Technical & Architectural Specification")
    run_tag.font.name = 'Segoe UI'
    run_tag.font.size = Pt(11)
    run_tag.font.italic = True
    run_tag.font.color.rgb = MUTED

    # Meta Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document Version", "1.0.0 (Production Release)"),
        ("Architecture", "Next.js 14 + NestJS 10 + GraphQL + Supabase"),
        ("Date", "October 2026"),
        ("Status", "Complete & Production Ready")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = SECONDARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Helper function for Section Headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(22)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = PRIMARY

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Segoe UI Semibold'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = SECONDARY

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Segoe UI Semibold'
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = PRIMARY

    def add_bullet(p_or_text, level=0):
        if isinstance(p_or_text, str):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(p_or_text)
            r.font.size = Pt(10)
            return p
        return p_or_text

    def add_styled_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header Row
        hdr_cells = table.rows[0].cells
        for idx, header_text in enumerate(headers):
            cell = hdr_cells[idx]
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]
            set_cell_background(cell, "2563EB")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(header_text)
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_values in enumerate(rows_data):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row_cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(51, 65, 85)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(8)

    # 1. Executive Summary
    add_h1("1. Executive Summary")
    p = doc.add_paragraph()
    p.add_run(
        "CanteenQ is an enterprise-grade campus canteen queue and food ordering ecosystem engineered to transform the traditional, congested dining hall into a streamlined digital experience. By replacing chaotic physical queues with real-time digital queuing, dynamic wait-time calculations, automated kitchen Kanban workflows, and cryptographic counter-pickup QR codes, CanteenQ eliminates counter congestion and delivers predictable meal ordering for colleges and universities."
    )
    p2 = doc.add_paragraph()
    p2.add_run(
        "Built on a cutting-edge decoupled architecture comprising a Next.js 14 App Router frontend hosted on Vercel, a high-throughput NestJS 10 backend exposing a Code-First GraphQL API alongside Socket.io WebSockets, and a Supabase Cloud PostgreSQL database managed via Prisma ORM, CanteenQ is designed for high concurrent throughput, instant response times, and automated CI/CD deployments via GitHub Actions."
    )

    # 2. Problem Statement & Background
    add_h1("2. Problem Statement & Motivation")
    add_h2("2.1 The Traditional Campus Dining Dilemma")
    doc.add_paragraph(
        "College canteens face severe surges during short recess and lunch intervals (e.g. 12:30 PM - 1:30 PM). During these peak windows, hundreds of students converge simultaneously on limited billing and food pickup counters, creating critical operational bottlenecks:"
    )
    add_bullet("Long Physical Queues: Students frequently spend 20 to 35 minutes standing in line just to place an order and pay, leaving minimal time to actually consume their meals.")
    add_bullet("Lack of Wait-Time Visibility: Before walking to the canteen, students have zero insight into how crowded the facility is, resulting in uncoordinated rushes.")
    add_bullet("Cash & Order Errors: Paper tokens, verbal order calling, and cash management cause frequent billing errors, lost tokens, and disputed meal handoffs.")
    add_bullet("Kitchen Chaos & Overload: Kitchen staff receive haphazard verbal spikes, leading to uneven preparation times, food stockouts, or excessive pre-cooking and food wastage.")
    add_bullet("Zero Operational Analytics: Campus administration lacks granular data on daily order volume, item popularity, hourly throughput, and bottleneck points.")

    add_h2("2.2 The CanteenQ Solution")
    doc.add_paragraph(
        "CanteenQ resolves these challenges through an end-to-end digital transformation:"
    )
    add_bullet("Pre-Ordering & Digital Queuing: Students order from their phones or laptops from any classroom or hostel, receiving an algorithmic wait time and dynamic queue position.")
    add_bullet("Live Campus Crowd Meter: An upfront digital load indicator (LOW / MODERATE / HIGH) lets students assess canteen occupancy before ordering.")
    add_bullet("Automated Kitchen Kanban: Kitchen chefs work from real-time digital boards, moving items through distinct preparation states (PLACED ➔ ACCEPTED ➔ PREPARING ➔ READY).")
    add_bullet("Cryptographic Counter Pickup: Food pickup is authenticated via dynamic, auto-generated QR codes, ensuring zero order theft or token mismatch.")
    add_bullet("Predictive AI Engine: Historical data forecasts peak rush windows and popular item demand, enabling optimal ingredient prep.")

    # 3. Technology Stack Deep-Dive
    add_h1("3. Technology Stack Deep-Dive")
    doc.add_paragraph(
        "The application leverages modern, production-grade open-source technologies selected for performance, type safety, scalability, and developer ergonomics:"
    )

    tech_headers = ["Layer", "Technology", "Version", "Key Responsibilities"]
    tech_rows = [
        ["Frontend UI", "Next.js (App Router)", "14.2.20", "Server/client components, routing, React 18, responsive layout"],
        ["Styling & Icons", "Tailwind CSS + Lucide", "3.4.16", "Modern dark-mode glassmorphic design system, lightweight SVG icons"],
        ["Data Visualization", "Recharts", "2.15.0", "Interactive revenue, hourly volume, and dish popularity graphs"],
        ["API Communication", "GraphQL (Apollo Server)", "12.2 / 4.11", "Code-First strongly typed queries, mutations, zero over-fetching"],
        ["Real-Time Gateway", "Socket.io", "4.8.1", "Bidirectional WebSocket pub-sub for live queue position and order stepper"],
        ["Backend Framework", "NestJS", "10.4.15", "Enterprise TypeScript framework, modular dependency injection, guards"],
        ["Cloud Database", "Supabase (PostgreSQL)", "16.x", "Cloud PostgreSQL with PgBouncer / Supavisor connection pooling"],
        ["ORM", "Prisma ORM", "5.22.0", "Type-safe database queries, declarative schema migrations, transactions"],
        ["Frontend Hosting", "Vercel", "Production", "Edge network distribution, continuous deployment, automated SSL"],
        ["Containerization", "Docker & Docker Compose", "20-Alpine", "Multi-stage reproducible container builds"],
        ["CI/CD Automation", "GitHub Actions", "v4 Workflows", "Automated linting, building, Jest tests, DB push, Vercel deploy"]
    ]
    add_styled_table(tech_headers, tech_rows, [Inches(1.3), Inches(1.5), Inches(0.9), Inches(2.8)])

    # 4. System Architecture & Data Flow
    add_h1("4. System Architecture & Operational Workflow")
    doc.add_paragraph(
        "CanteenQ employs a modern decoupled client-server architecture. The Next.js frontend acts as the user interface layer, communicating with the NestJS backend via a unified GraphQL endpoint (/graphql) for state mutations and queries, while maintaining an active WebSocket connection for instantaneous broadcast events."
    )

    add_h2("4.1 Core Portals by User Role")
    add_h3("1. Student Portal (/menu, /cart, /checkout, /order/[id], /orders)")
    doc.add_paragraph(
        "Designed for rapid ordering and real-time visibility from mobile devices and laptops:"
    )
    add_bullet("Live Crowd Density: An intuitive circular and pill meter displays current dining hall load percentage (0-40% LOW, 41-70% MODERATE, 71-100% HIGH), queue length, and estimated wait minutes.")
    add_bullet("Searchable Interactive Catalog: Filters by category (Breakfast, Meals, Snacks, Beverages, Fast Food, Desserts), search query, and Veg/Non-Veg toggle.")
    add_bullet("Cart & Checkout: Persistent local shopping cart with real-time tax calculation, add/remove controls, and payment method selection (UPI, Card, Net Banking, or Instant Test Gateway).")
    add_bullet("Live Order Stepper: Real-time visual progress tracker moving across 6 states: PLACED ➔ PAID ➔ ACCEPTED ➔ PREPARING ➔ READY ➔ COMPLETED.")
    add_bullet("Dynamic Queue Counter: Shows exact count of orders ahead in the kitchen queue, automatically counting down as prior orders finish.")
    add_bullet("Cryptographic Pickup QR Code: Automatically renders upon reaching the 'READY' state for counter verification.")

    add_h3("2. Kitchen Staff / Chef Portal (/staff)")
    doc.add_paragraph(
        "Provides kitchen leads and cooking stations with complete operational control:"
    )
    add_bullet("Kitchen Kanban Board: Orders automatically enter 4 columns: New Orders (Placed/Accepted) ➔ In Preparation ➔ Ready for Counter Pickup ➔ Completed Today.")
    add_bullet("One-Click State Transitions: Chefs click single buttons ('Start Preparing', 'Mark Ready', 'Complete') which instantly broadcast WebSocket events to the student's active screen.")
    add_bullet("Live Stock & Inventory Manager: Instant toggling of dishes between Available and Sold Out, with real-time stock number and prep time adjustments.")

    add_h3("3. Campus Admin & Analytics Portal (/admin)")
    doc.add_paragraph(
        "Delivers executive oversight, business analytics, and predictive planning tools:"
    )
    add_bullet("Real-Time KPI Cards: Today's revenue, order count, average waiting time, kitchen turnaround velocity, and order rejection rates.")
    add_bullet("Visual Recharts Charts: Hourly order distribution, revenue breakdown, and top-performing dishes.")
    add_bullet("AI Prediction Engine: Forecasts peak rush windows (e.g. 8:30 AM - 9:30 AM, 12:30 PM - 2:00 PM), expected order volume for the next 4 hours, and high-demand food item prep requirements.")

    # 5. Database Architecture & Schema
    add_h1("5. Database Architecture & Schema Design")
    doc.add_paragraph(
        "The data layer is modeled using Prisma ORM connected to Supabase Cloud PostgreSQL. The database consists of 14 relational models supporting referential integrity, foreign key constraints, and cascading rules:"
    )

    db_headers = ["Model", "Primary Key", "Key Attributes", "Relational Connections"]
    db_rows = [
        ["User", "id (CUID)", "name, email, role, passwordHash, isActive", "1:1 Student, 1:1 Staff, 1:N Notification"],
        ["Student", "id (CUID)", "studentId, department, semester, className", "1:1 User, 1:1 Cart, 1:N Order"],
        ["Staff", "id (CUID)", "employeeId, designation, shift", "1:1 User"],
        ["Canteen", "id (CUID)", "name, location, openingTime, closingTime, maxCapacity", "Stand-alone master entity"],
        ["Category", "id (CUID)", "name (unique), description", "1:N FoodItem"],
        ["FoodItem", "id (CUID)", "name, price, preparationTime, stockQuantity, isAvailable, isVeg", "N:1 Category, 1:N CartItem, 1:N OrderItem"],
        ["Cart", "id (CUID)", "studentId (unique), totalAmount", "1:1 Student, 1:N CartItem"],
        ["CartItem", "id (CUID)", "quantity, unitPrice, subtotal", "N:1 Cart, N:1 FoodItem"],
        ["Order", "id (CUID)", "orderNumber, status, paymentStatus, totalAmount, waitingTime", "N:1 Student, 1:N OrderItem, 1:1 QueueToken, 1:1 Payment"],
        ["OrderItem", "id (CUID)", "quantity, unitPrice, subtotal", "N:1 Order, N:1 FoodItem"],
        ["QueueToken", "id (CUID)", "tokenNumber (SQ-xxx), position, status, waitTime", "1:1 Order"],
        ["Payment", "id (CUID)", "transactionId, provider, amount, currency, status", "1:1 Order"],
        ["Notification", "id (CUID)", "title, message, type, isRead, createdAt", "N:1 User"],
        ["QueueSnapshot", "id (CUID)", "timestamp, queueLength, activeOrders, crowdPercentage", "Time-series logging for predictions"]
    ]
    add_styled_table(db_headers, db_rows, [Inches(1.2), Inches(1.1), Inches(2.2), Inches(2.0)])

    # 6. GraphQL API Specification
    add_h1("6. GraphQL API Specification")
    doc.add_paragraph(
        "CanteenQ uses a Code-First GraphQL schema implemented via @nestjs/graphql and Apollo Server on http://localhost:4000/graphql. All client communication uses strongly typed queries and mutations:"
    )

    gql_headers = ["Domain", "Operation", "Type", "Access / Guard", "Description"]
    gql_rows = [
        ["Auth", "login(email, password)", "Mutation", "Public", "Authenticates user and returns JWT accessToken + profile"],
        ["Auth", "register(input)", "Mutation", "Public", "Registers student account with profile and empty cart"],
        ["Menu", "getMenu(filter)", "Query", "Public", "Fetches dishes filtered by category, veg, search, and stock"],
        ["Menu", "getCategories", "Query", "Public", "Fetches all food categories with descriptions"],
        ["Menu", "updateFoodAvailability", "Mutation", "Staff / Admin", "Updates item stock quantity and availability toggle"],
        ["Queue", "getCrowdDensity", "Query", "Public", "Calculates live crowd percentage, wait minutes, active orders"],
        ["Orders", "createOrder(paymentMethod)", "Mutation", "Student", "Converts cart to order, reserves stock, creates SQ token"],
        ["Orders", "getMyOrders", "Query", "Student", "Returns past and active orders for logged-in student"],
        ["Orders", "acceptOrder(id)", "Mutation", "Staff", "Moves order from PLACED to ACCEPTED; notifies student"],
        ["Orders", "startPreparing(id)", "Mutation", "Staff", "Moves order to PREPARING status"],
        ["Orders", "markReady(id)", "Mutation", "Staff", "Marks READY, emits pickup QR code to student screen"],
        ["Orders", "completeOrder(id)", "Mutation", "Staff", "Finalizes order pickup, releases queue token"],
        ["Payment", "createPayment / verifyPayment", "Mutation", "Student", "Processes and validates payment status"],
        ["Analytics", "getAnalytics", "Query", "Admin", "Aggregates revenue, turnaround times, and dish popularity"],
        ["Predictions", "getPredictions", "Query", "Admin", "Forecasts peak rush hours and ingredient demands"]
    ]
    add_styled_table(gql_headers, gql_rows, [Inches(0.9), Inches(1.8), Inches(0.8), Inches(1.0), Inches(2.0)])

    # 7. Deployment & CI/CD
    add_h1("7. Deployment & CI/CD Pipeline")
    doc.add_paragraph(
        "The application is engineered for production deployment with continuous integration and continuous deployment (CI/CD):"
    )
    add_bullet("Frontend (Vercel): Next.js App Router deployed with zero-config edge rendering, security headers (nosniff, frame-options DENY, strict-origin), and automatic branch previews.")
    add_bullet("Backend: Multi-stage Docker containerization running on Node 20 Alpine, deployable to Railway, Render, Fly.io, or AWS ECS/EC2.")
    add_bullet("Database (Supabase): Dual connection string strategy utilizing pooled connections (port 6543) for runtime queries and direct connections (port 5432) for Prisma schema pushes.")
    add_bullet("Automated GitHub Actions (.github/workflows/ci-cd.yml): Runs 5 automated stages on push/PR to main: Backend Build & Tests, Frontend Production Build, Docker Validation, Supabase Schema Push, and Automated Vercel Deployment.")

    # 8. Future Scope & Roadmap
    add_h1("8. Future Scope & Evolutionary Roadmap")
    doc.add_paragraph(
        "CanteenQ provides a scalable foundation that can be extended with next-generation smart campus features:"
    )
    add_bullet("Campus RFID / Smart Card Integration: Tap student university ID cards at dedicated counter terminals to debit existing campus meal allowances or student wallet balances.")
    add_bullet("Computer Vision Crowd Meter: Mount overhead cameras at canteen entrances running lightweight YOLO/OpenCV models to count real-world dining hall occupancy and validate virtual queue density.")
    add_bullet("Dedicated Kitchen Display System (KDS): Hardware touch terminals with waterproof bump bars and wireless ESC/POS thermal ticket printers for kitchen prep stations.")
    add_bullet("Automated Ingredient & Inventory Replenishment: Connect food item stock depletion directly with supplier purchase orders to automatically re-order raw ingredients before stockouts occur.")
    add_bullet("Multi-Canteen & Multi-Vendor Expansion: Support universities with multiple dining halls, cafeterias, coffee kiosks, and night canteens within a single unified app.")
    add_bullet("Personalized Nutritional & Allergen Profiling: Provide students with calorie counters, macro breakdowns (protein, carbs, fats), and automated allergen alerts (gluten-free, nut-free, dairy-free).")

    # 9. Conclusion
    add_h1("9. Conclusion")
    doc.add_paragraph(
        "CanteenQ delivers an elegant, high-impact solution to one of the most persistent operational pain points in campus life: long, unpredictable canteen queues. By harmonizing Next.js 14, NestJS 10, GraphQL, WebSockets, Supabase Cloud PostgreSQL, and Prisma ORM into a cohesive, fault-tolerant platform, CanteenQ elevates campus dining into an efficient, predictable, and modern student-centric experience."
    )

    # Save Document
    output_path = os.path.abspath("CanteenQ_Project_Documentation.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
